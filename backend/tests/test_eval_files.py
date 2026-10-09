"""The stress-test files in data/eval, scored without any outside service.

Each file was built with known problems in it (data/eval/make_files.py), so
these tests hold the rules and the column detection to an answer key on files
shaped nothing like the samples the app grew up on. The half that needs a model
or a lookup service is run on demand with data/eval/evaluate.py.
"""

import io
import sys
from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook

EVAL = Path(__file__).resolve().parents[2] / "data" / "eval"
sys.path.insert(0, str(EVAL))

from evaluate import Score, StandInSymbols, score_offline  # noqa: E402
from make_files import Key, build  # noqa: E402

from app.config import Settings  # noqa: E402
from app.connectors.build import build_runtime  # noqa: E402
from app.connectors.cache import LookupCache  # noqa: E402
from app.store import DatasetStore  # noqa: E402

FILES = {key.filename: (data, key) for data, key in build()}


@pytest.fixture(scope="module")
def scores() -> Iterator[dict[str, Score]]:
    settings = Settings()
    runtime = build_runtime(settings, cache=LookupCache())
    runtime.symbols = StandInSymbols()
    store = DatasetStore(settings)
    found = {
        name: score_offline(data, key, store, runtime)[0] for name, (data, key) in FILES.items()
    }
    yield found
    store.close_all()


def key_of(name: str) -> Key:
    return FILES[name][1]


def _contents(name: str, data: bytes) -> object:
    """What a file holds. An Excel file also carries the time it was written."""
    if not name.endswith(".xlsx"):
        return data
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    return {
        sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)] for sheet in workbook
    }


def test_the_same_seed_writes_the_same_files() -> None:
    again = {key.filename: _contents(key.filename, data) for data, key in build()}

    assert again == {name: _contents(name, data) for name, (data, _) in FILES.items()}


@pytest.mark.parametrize("name", sorted(FILES))
def test_no_file_breaks_the_review(name: str, scores: dict[str, Score]) -> None:
    score = scores[name]

    assert score.rows > 0
    assert score.header_ok
    assert score.errors == []


@pytest.mark.parametrize("name", sorted(FILES))
def test_every_problem_put_into_a_file_gets_a_card(name: str, scores: dict[str, Score]) -> None:
    assert scores[name].problems_missed == []


@pytest.mark.parametrize("name", sorted(FILES))
def test_look_alike_columns_are_never_offered_to_a_connector(
    name: str, scores: dict[str, Score]
) -> None:
    # Batch numbers, round quantities, years, state codes, product lines called
    # "Garden Group": none is a postal code, a company or a stock symbol.
    assert scores[name].offers_on_traps == []


@pytest.mark.parametrize("name", sorted(FILES))
def test_columns_that_can_be_looked_up_are_offered(name: str, scores: dict[str, Score]) -> None:
    assert scores[name].offers_missed == []


@pytest.mark.parametrize("name", sorted(FILES))
def test_cleaned_totals_and_row_counts_match_the_answer_key(
    name: str, scores: dict[str, Score]
) -> None:
    score = scores[name]

    assert [line for line in score.totals.values() if not line.startswith("ok")] == []
    assert score.rows_after == "" or score.rows_after.startswith("ok")


def test_a_tidy_file_gets_nothing_to_decide(scores: dict[str, Score]) -> None:
    # Reading dates and numbers as what they are is fine. A question or a
    # suggestion on a file with nothing wrong would be noise.
    cards = scores["already_clean.csv"].issues

    assert [card for card in cards if not card.startswith("safe:")] == []


def test_column_names_under_title_lines_are_found(scores: dict[str, Score]) -> None:
    # Three lines sit above the survey's header row. Before they were skipped
    # its columns were "column_3" and its cards named nothing a reader knew.
    score = scores["survey.xlsx"]

    assert score.errors == []
    assert score.rows == 200
    assert any('"Would recommend"' in card for card in score.issues)


def test_european_prices_and_postal_codes_keep_their_meaning(scores: dict[str, Score]) -> None:
    score = scores["orders_eu.csv"]

    # "1.234,56 €" adds up, and "01067" is not turned into the number 1067.
    assert score.totals["Einzelpreis"].startswith("ok")
    assert not any('Read "PLZ" as numbers' in card for card in score.issues)
    assert ("postal_geocode", "PLZ") in score.offers


def test_a_wide_file_is_reviewed_in_reasonable_time(scores: dict[str, Score]) -> None:
    # 70 columns raised 129 cards and took nine seconds to rank before impact
    # scoring had a ceiling. Generous, so a slow machine doesn't fail it.
    assert scores["inventory_traps.csv"].seconds["find problems"] < 8

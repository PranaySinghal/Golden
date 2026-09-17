import datetime as dt
import io
import math
import zipfile
from collections.abc import Iterator

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook

from app.errors import SheetNotFoundError, UnreadableFileError
from app.ingest.models import IngestOptions, SourceFormat, SourceTable

# Floats above this can't be shown exactly as whole numbers, so they keep float formatting.
_MAX_EXACT_WHOLE = 2**53


class XlsxAdapter:
    format = SourceFormat.XLSX

    def read(self, data: bytes, options: IngestOptions) -> SourceTable:
        workbook = _open_workbook(data)
        try:
            sheet_names = [ws.title for ws in workbook.worksheets]
            if options.sheet is not None:
                if options.sheet not in sheet_names:
                    raise SheetNotFoundError(options.sheet, sheet_names)
                sheet_name = options.sheet
            else:
                sheet_name = _first_visible_sheet(workbook)
        except BaseException:
            workbook.close()
            raise

        return SourceTable(
            rows=_read_rows(workbook, sheet_name),
            sheet_name=sheet_name,
            available_sheets=sheet_names,
        )


def cell_to_text(value: object) -> str | None:
    """Render a cell the way a spreadsheet user would see it, as plain text.

    Excel stores every number as a float, so whole numbers drop the trailing
    ".0" that nobody typed.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer() and abs(value) < _MAX_EXACT_WHOLE:
            return str(int(value))
        return repr(value)
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    text = str(value)
    return text if text != "" else None


def _open_workbook(data: bytes) -> Workbook:
    try:
        # data_only returns the last computed value of formulas instead of the formula text.
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise UnreadableFileError(
            "We couldn't open this Excel file. It may be damaged or password protected. "
            "Try saving it again from Excel."
        ) from exc


def _first_visible_sheet(workbook: Workbook) -> str:
    for ws in workbook.worksheets:
        if getattr(ws, "sheet_state", "visible") == "visible":
            return ws.title
    raise UnreadableFileError("This Excel file has no visible sheets to read.")


def _read_rows(workbook: Workbook, sheet_name: str) -> Iterator[list[str | None]]:
    try:
        ws = workbook[sheet_name]
        # Some tools write a wrong sheet size; resetting makes openpyxl read what's there.
        reset = getattr(ws, "reset_dimensions", None)
        if callable(reset):
            reset()
        for row in ws.iter_rows(values_only=True):
            yield [cell_to_text(value) for value in row]
    finally:
        workbook.close()
